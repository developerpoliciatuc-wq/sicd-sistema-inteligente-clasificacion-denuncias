from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)


_HEADERS = [
    "timestamp",
    "fecha_denuncia",
    "region",
    "comisaria",
    "tipo_delito",
    "modalidad_delito",
    "direccion_hecho",
    "latitud",
    "longitud",
    "precision_geo",
    "archivo_origen",
    "archivo_destino",
    "status",
    "motivo",
]


def append_stats(xlsx_path: Path, row: dict) -> None:
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)

    if xlsx_path.exists():
        wb = load_workbook(xlsx_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(_HEADERS)

    ws.append([row.get(h) for h in _HEADERS])
    wb.save(xlsx_path)
